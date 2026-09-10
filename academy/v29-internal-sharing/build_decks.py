"""Editable bilingual decks and scene-rendered PDF previews for manuscript v29.

PPTX objects remain editable. Previews use the same measured layout; they are
not a PowerPoint/LibreOffice rendering. Fonts are Noto Sans CJK SC.
Run: python build_decks.py
"""
from pathlib import Path
import json
import re
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.xmlchemy import OxmlElement
from reportlab.pdfgen.canvas import Canvas
from reportlab.lib.utils import ImageReader

ROOT = Path(__file__).resolve().parent
W, H, SCALE = 13.333333, 7.5, 120
INK, MUTED, BLUE, TEAL, PALE, GOLD = '152536', '596A78', '154C79', '087F8C', 'EEF4F7', '9A6517'
FONT = '/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc'
BOLD = '/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc'
PAPER = 'https://arxiv.org/abs/2608.27086'
LOCAL = '../agentic-runtime-preprint/paper_source/main.tex'
SOURCES = [
 ('Prime Agent', 'https://arxiv.org/html/2608.23552v1'),
 ('AutoResearch', 'https://arxiv.org/html/2608.17906v1'),
 ('Apodex 1.1', 'https://arxiv.org/html/2608.23283v1'),
 ('VoiceMem', 'https://arxiv.org/html/2608.26005v1'),
 ('NeoHorse-1', 'https://arxiv.org/html/2609.08183v1'),
 ('Beyond Top-K', 'https://arxiv.org/html/2608.06305v1'),
]


class Slide:
    def __init__(self, deck, lang, title, subtitle='', sources=None, notes=''):
        self.slide = deck.slides.add_slide(deck.slide_layouts[6])
        self.lang = lang
        self.image = Image.new('RGB', (1600, 900), 'white')
        self.draw = ImageDraw.Draw(self.image)
        self.texts = []
        self.rect(.45, .37, .07, .48, TEAL)
        self.text(.68, .32, 12.05, .72, title, 28, True)
        if subtitle:
            self.text(.68, 1.08, 12, .55, subtitle, 15, color=MUTED)
        self.line(.68, 6.88, 12.66, 6.88, 'CBD5DD')
        self.text(.68, 7.02, 9.2, .2,
                  '内部分享 · 研究稿 v29 · 2026-09-10' if lang=='zh' else
                  'Internal sharing · Research manuscript v29 · 2026-09-10', 9, color=MUTED)
        self.text(12.0, 7.01, .65, .25, f'{len(deck.slides):02d}', 10, color=MUTED)
        self.notes = notes
        self.sources = sources or []
        for i, (label, url) in enumerate(self.sources[:3]):
            self.text(.68 + i*4.05, 6.6, 3.95, .22, label, 9, color=MUTED, link=url)
        self.slide.notes_slide.notes_text_frame.text = notes + '\n\n' + '\n'.join(f'{a}: {b}' for a,b in self.sources)

    def rect(self,x,y,w,h,fill=PALE):
        shape=self.slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
        shape.fill.solid(); shape.fill.fore_color.rgb=RGBColor.from_string(fill)
        shape.line.fill.background()
        self.draw.rectangle((x*SCALE,y*SCALE,(x+w)*SCALE,(y+h)*SCALE),fill='#'+fill)
        return shape

    def line(self,x1,y1,x2,y2,color=TEAL,width=1.5):
        shape=self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
        shape.line.color.rgb=RGBColor.from_string(color); shape.line.width=Pt(width)
        self.draw.line((x1*SCALE,y1*SCALE,x2*SCALE,y2*SCALE),fill='#'+color,width=max(1,round(width*SCALE/72)))

    def text(self,x,y,w,h,txt,size=20,bold=False,color=INK,link=None):
        # Measure line breaks once, so PPT and preview use identical wrapping.
        font=ImageFont.truetype(BOLD if bold else FONT,round(size*SCALE/72))
        lines=[]
        for para in str(txt).split('\n'):
            tokens=list(para) if self.lang=='zh' else re.findall(r'\S+\s*',para)
            line=''
            for token in tokens:
                if line and self.draw.textlength(line+token,font=font)>w*SCALE-4:
                    lines.append(line.rstrip()); line=token.lstrip()
                else: line+=token
            lines.append(line.rstrip())
        pitch=size/72*1.32
        if len(lines)*pitch > h+.035:
            if size <= 14:
                raise ValueError(f'Text overflow slide: {txt[:55]} ({len(lines)*pitch:.2f}>{h})')
            return self.text(x,y,w,h,txt,size-1,bold,color,link)
        assert x>=0 and y>=0 and x+w<=W+.01 and y+h<=H+.01
        box=self.slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
        tf=box.text_frame; tf.clear(); tf.word_wrap=False; tf.auto_size=MSO_AUTO_SIZE.NONE
        tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
        for i,line in enumerate(lines):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
            p.space_before=Pt(0); p.space_after=Pt(0); p.line_spacing=Pt(size*1.32)
            r=p.add_run();r.text=line;r.font.name='Noto Sans CJK SC';r.font.size=Pt(size)
            r.font.bold=bold;r.font.color.rgb=RGBColor.from_string(color)
            ea=OxmlElement('a:ea');ea.set('typeface','Noto Sans CJK SC');r._r.get_or_add_rPr().append(ea)
            if link:r.hyperlink.address=link
            self.draw.text((x*SCALE,(y+i*pitch)*SCALE),line,font=font,fill='#'+color,anchor='lt')
        self.texts.append(str(txt))

    def panel(self,x,y,w,h,title,body,color=TEAL,size=19):
        self.rect(x,y,w,h)
        self.rect(x,y,.05,h,color)
        self.text(x+.2,y+.18,w-.4,.52,title,21,True,color)
        self.text(x+.2,y+.92,w-.4,h-1.08,body,size)

    def takeaway(self,txt):
        self.text(.75,5.86,11.9,.56,txt,19,True,color=BLUE)


def build(lang):
    zh=lang=='zh'
    def L(a,b):return a if zh else b
    deck=Presentation();deck.slide_width=Inches(W);deck.slide_height=Inches(H)
    deck.core_properties.title=L('契约驱动的 Agentic Runtime：v29 内部分享','Contract-Centered Agentic Runtimes: v29 Internal Sharing')
    deck.core_properties.subject='Data Wiki, Theme Wiki, 5W1H+Which, executable contracts and evidence boundaries'
    deck.core_properties.author='Yaxiao Liu; Pengbo Liu; Yiwen Liu; Yihua Guan; Zhenghe Hou; Jiaxing Song'
    slides=[]
    def S(title,sub='',refs=None,note=''):
        s=Slide(deck,lang,title,sub,refs or [('v29 manuscript / 论文源码',LOCAL)],note)
        slides.append(s);return s

    s=S(L('契约驱动的 Agentic Runtime','Contract-Centered Agentic Runtimes'),
        L('从三层责任架构，到可执行的 5W1H 数据使用契约','From responsibility boundaries to executable 5W1H data-use contracts'),
        [('Existing arXiv / 已发表前版',PAPER)],
        L('建议用25分钟讲解，另留10分钟讨论。先说明这是内部版本v29，已有arXiv论文是前版。本次真正新增的是可执行数据边界与有限验证，不是完整系统或顶会录用结果。作者顺序沿用论文。','Allow 25 minutes plus discussion. Internal v29 revises the existing arXiv manuscript. The new contribution is an executable data boundary with finite validation, not a completed production system or a conference acceptance.'))
    s.text(.8,2.05,11.8,1.3,L('能力可以迭代，容量可以扩展。\n数据的身份、授权与证据必须始终可检查。','Capabilities evolve. Capacity grows.\nData identity, authority and evidence must remain checkable.'),30,True)
    s.rect(.8,3.85,11.75,.8,PALE)
    s.text(1.05,4.04,11.25,.4,L('研究状态：机制与合成验证已落地；真实任务和规模实验待完成','Status: executable mechanism and synthetic checks; real-task and scaling studies remain open'),18,color=BLUE)
    s.text(.8,5.12,11.75,.85,'Yaxiao Liu · Pengbo Liu · Yiwen Liu · Yihua Guan · Zhenghe Hou · Jiaxing Song\nPwC China AI Center / Tsinghua University',14,color=MUTED)

    s=S(L('一个任务成功，不代表一个系统可管理','Task success does not establish system manageability'),
        L('企业任务同时跨越业务、模型、基础设施和数据治理边界','Enterprise work crosses business, model, infrastructure and data-governance boundaries'),note=L('用一份季度经营报告引入。最后数字看起来正确，仍可能使用错期间、越权数据或旧文件。我们要同时问业务任务是否完成，以及运行路径是否满足约束。','Introduce a quarterly report. A plausible final number can still rely on the wrong period, unauthorized evidence or a stale artifact. Ask both whether the task was completed and whether its execution met the contract.'))
    s.panel(.75,1.95,3.8,3.3,L('业务问题','Business'),L('结果是否有用？\n所有必需字段是否正确？\n是否需要人工返工？','Is the result useful?\nAre all required fields correct?\nHow much human rework remains?'))
    s.panel(4.76,1.95,3.8,3.3,L('系统问题','Runtime'),L('容量变化是否改变行为？\n失败能否恢复？\n控制开销是否可接受？','Does capacity change behavior?\nCan failures be recovered?\nAre enforcement costs acceptable?'))
    s.panel(8.77,1.95,3.8,3.3,L('数据问题','Data'),L('来源版本是否正确？\n此刻是否仍有授权？\n结论能否追到原始证据？','Is the source version correct?\nIs access still authorized now?\nCan claims be traced to evidence?'))
    s.takeaway(L('业务证据与运行时证据必须分别给出。','Business evidence and runtime evidence must be reported separately.'))

    s=S(L('v29 将设计主张推进到可执行边界','v29 makes the data boundary executable'),note=L('强调三项贡献的证据等级不同。责任架构是设计，契约参考模型可执行，P1是仍待验证的经验假设。不要把它们统称为已验证架构。','The three contributions have different evidence levels. The architecture is a design; the data reference model executes; P1 remains an empirical hypothesis. Do not call the complete architecture validated.'))
    for i,(a,b,c) in enumerate([
      (L('01 责任契约','01 Responsibility contracts'),L('Skill / Harness / Scaffold\n外部数据独立治理','Skill / Harness / Scaffold\nIndependent external data governance'),L('架构设计','Design')),
      (L('02 数据使用边界','02 Data-use boundary'),L('5W1H+Which 谓词\n执行期重验证与证据绑定','5W1H+Which predicates\nRevalidation and evidence binding'),L('可运行的有限参考模型','Executable, bounded reference model')),
      (L('03 可证伪评估','03 Falsifiable evaluation'),L('能力×容量交叉实验\n质量、违规与总成本','Capability × capacity experiment\nQuality, violations and total cost'),L('确认性实验尚未执行','Confirmatory study not yet run'))]):
        x=.75+i*4.01;s.panel(x,1.9,3.8,3.25,a,b);s.text(x+.15,5.32,3.55,.48,c,15,True,color=BLUE)

    s=S(L('四个责任对象，三个运行时层','Four responsibility objects; three runtime layers'),
        L('数据基座位于运行栈外，拥有独立的语义和治理权','The data substrate sits outside the stack with independent semantic ownership'),note=L('Skill声明业务能力，Harness进行准入和绑定，Scaffold提供执行边界与资源。外部数据基座不调度计算，也不被某个运行层私有化。这是责任模型，不要求企业必须设四个团队。','Skill declares business capability. Harness admits and binds it. Scaffold supplies execution boundaries and resources. The external substrate governs data semantics. This is a responsibility model, not a mandatory four-team organization chart.'))
    for y,title,body in [(1.85,'Skill',L('可复用能力、输入输出、允许效果、版本','Reusable capability, I/O, allowed effects and version')),(3.03,'Harness',L('准入、图编排、策略检查、资源绑定、轨迹','Admission, orchestration, policy, binding and traces')),(4.21,'Scaffold',L('计算、隔离、身份、地域、恢复与资源核算','Compute, isolation, identity, locality, recovery and accounting'))]:
        s.rect(.8,y,7.3,.98);s.text(1.02,y+.12,2.0,.4,title,22,True,color=TEAL);s.text(3.1,y+.17,4.75,.66,body,16)
    s.panel(8.65,1.85,3.9,3.34,L('外部数据基座','External data substrate'),L('来源与快照\n语义、授权、血缘\nData Wiki / Theme / IR','Sources and snapshots\nSemantics, access, lineage\nData Wiki / Theme / IR'),size=18)
    s.line(8.12,3.52,8.6,3.52);s.takeaway(L('职责可以由同一团队实现，但契约边界仍需可检查。','One team may implement several objects; their contracts remain distinct.'))

    s=S(L('P1：能力与容量是否能有条件分离？','P1: can capability and capacity be separated?'),
        L('研究的是交互效应、语义稳定性与控制成本','Test interaction effects, semantic stability and enforcement cost'),note=L('以两种能力配置、两种容量配置解释2×2实验。扩容改善时延是预期主效应，不是P1的证明。需要检验改变能力的效果是否因容量改变而变化，同时语义不退化、开销不超预算。阈值必须由决策确定。','Explain the 2×2 experiment. A capacity main effect is expected and does not establish P1. Test whether the capability effect changes across capacity, while semantics remain non-inferior and control cost stays within budget. Margins must be decision-derived.'))
    s.text(.9,1.9,5.7,.4,L('同一工作负载、模型、策略和源快照','Same workload, model, policy and source snapshot'),17,True)
    for r in range(2):
        for c in range(2):
            x=.9+c*2.8;y=2.65+r*1.35;s.rect(x,y,2.6,1.13,PALE if r==0 else 'E3EFF0')
            s.text(x+.15,y+.18,2.3,.7,f'c{r} × s{c}\nR / Q / E',20,True,color=BLUE)
    s.panel(7.0,1.9,5.5,3.88,L('每个单元同时观测','Measure every cell'),L('R：时延、吞吐、超时\nQ：所有要求均满足的比例\nE：准入、检查、证据处理开销','R: latency, throughput, timeouts\nQ: every-requirement satisfaction\nE: admission, checks and evidence cost'))
    s.takeaway(L('P1 仍未被本研究的真实运行时实验验证。','P1 remains untested in a real runtime experiment in this study.'))

    s=S(L('最小反例：列名相同，业务意义不同','A minimal counterexample: same columns, different meaning'),
        L('示意场景；不是实测数据','Illustrative scenario; not measured data'),note=L('不要讲成一次SQL join就能发现的普通类型错误。两张表在语法和字段名上都可能一致，但单位、范围和期间不同。我们要求这些语义进入契约，而不是依赖模型临场猜测。','This is not merely an SQL type error. Schemas and field names can agree while units, entity scope and periods differ. The proposal makes those constraints explicit rather than relying on model inference at query time.'))
    s.panel(.8,1.95,5.55,2.55,L('来源 A：集团季度收入','Source A: group revenue'),L('字段：product_id, revenue\n口径：集团合并 / USD / 2026 Q2','Fields: product_id, revenue\nScope: consolidated / USD / 2026 Q2'))
    s.panel(6.6,1.95,5.9,2.55,L('来源 B：子公司收入','Source B: subsidiary revenue'),L('字段：product_id, revenue\n口径：本地公司 / EUR / 2026 Q1','Fields: product_id, revenue\nScope: local entity / EUR / 2026 Q1'),color=GOLD)
    s.text(.95,4.95,11.5,.68,L('检索相关、连接可执行，仍不代表可用于当前报告。','Relevant retrieval and executable joins do not establish admissibility.'),24,True,color=BLUE)

    s=S(L('Data Wiki、Theme Wiki 与 IR 各回答什么','Data Wiki, Theme Wiki and IR answer different questions'),note=L('Data Wiki面向数据本身，Theme面向结果要求，IR记录本任务怎样使用这些数据。摘要可以帮助发现，不能替代原始来源。IR在这里是Intermediate Relation，执行票据是从关系中解析出来的另一个对象。','Data Wiki describes sources. Theme Wiki specifies outputs. IR declares how a task may connect them. Summaries aid discovery, but canonical sources remain authoritative. IR means Intermediate Relation; the execution ticket is a separate resolved object.'))
    for x,title,body in [(.8,'Data Wiki',L('这是什么数据？\n来源、版本、字段、策略\n可调用的读／抽取算子','What is this data?\nSource, version, fields, policy\nRead and extraction operators')),(4.85,'IR',L('这次任务怎样使用？\n5W1H+Which\n证据、兼容性、回退','How may this task use it?\n5W1H+Which\nEvidence, compatibility, fallback')),(8.9,'Theme Wiki',L('产出必须满足什么？\n字段、单位、期间\n质量要求与验证器','What must the output satisfy?\nFields, units, periods\nQuality criteria and verifier'))]:
        s.panel(x,2.05,3.6,3.35,title,body,size=18)
    s.line(4.4,3.7,4.83,3.7);s.line(8.45,3.7,8.88,3.7)
    s.takeaway(L('自然语言描述是元数据；原始来源与可信策略提供权威。','Descriptions are metadata; sources and trusted policy remain authoritative.'))

    s=S(L('让 5W1H+Which 成为可检查条件','Turn 5W1H+Which into checkable conditions'),note=L('这些问句不是七项独立创新。技术工作在于把字段连接到可信输入、执行检查和明确故障。Why只是批准的整合理由，不是模型生成的因果证明。当前参考模型只实现简化检查，不实现多源join。','The seven question words are not seven inventions. The technical work links them to trusted inputs, runtime checks and explicit failures. Why is an approved rationale, not a model-generated causal proof. The reference model implements simplified checks, not multi-source joins.'))
    rows=[('When',L('业务期间、有效期、当前时刻','Business period, expiry, current time'),L('排队后过期','Expiry while queued')),('Where',L('访问域与交付边界','Access domain and delivery boundary'),L('内部数据外流','Internal data released externally')),('Who',L('可信身份与当前授权','Trusted identity and current access'),L('准入后撤权','Revocation after admission')),('What',L('单位、实体范围、业务口径','Units, entity scope, business meaning'),L('同名异义','Same name, different meaning')),('Why',L('批准的整合理由标识','Approved integration rationale'),L('编造连接理由','Invented integration rationale')),('How',L('允许的算子和预算','Allowed operations and budget'),L('读请求变成写','Read request becomes write')),('Which',L('关系、版本、键、血缘','Relations, versions, keys, lineage'),L('错误关联或快照','Wrong relationship or snapshot'))]
    for i,(a,b,c) in enumerate(rows):
        y=1.85+i*.52
        if i%2==0:s.rect(.8,y,11.7,.49,PALE)
        s.text(1,y+.06,1.45,.4,a,17,True,color=TEAL);s.text(2.5,y+.06,6.25,.4,b,16);s.text(9,y+.06,3.1,.4,c,15,color=MUTED)
    s.takeaway(L('字段存在不等于约束被执行。','A populated field does not mean its constraint is enforced.'))

    s=S(L('从候选检索，到执行票据，再到证据','From candidates to execution tickets to evidence'),note=L('分清发现和授权：搜索结果只是候选，不能直接授予工具或源权限。执行前重查版本与策略，返回字节后绑定证据。失败进入typed rejection，回退只能选择已允许路径。','Discovery is not authorization. Search returns candidates; it cannot grant tools or source privileges. Recheck versions and policy immediately before execution, then bind returned evidence. Fallback can only use already permitted paths.'))
    steps=[L('发现候选','Discover'),L('解析契约','Resolve'),L('绑定票据','Bind ticket'),L('执行期重查','Revalidate'),L('读取与证据','Read + bind')]
    bodies=[L('排序不授予权限','Ranking grants no authority'),L('可信策略与版本','Trusted policy and versions'),L('身份、epoch、预算','Identity, epoch, budget'),L('过期／撤销／源变化','Expiry, revocation, source change'),L('字节、摘要、来源','Bytes, digest, source')]
    for i,(a,b) in enumerate(zip(steps,bodies)):
        x=.8+i*2.36;s.panel(x,2.2,2.18,2.72,a,b,size=17)
        if i<4:s.line(x+2.18,3.57,x+2.35,3.57)
    s.takeaway(L('失败有类型：STALE / UNAUTHORIZED / SOURCE_CHANGED / REVALIDATE','Typed failures: STALE / UNAUTHORIZED / SOURCE_CHANGED / REVALIDATE'))

    s=S(L('撤权发生在排队期间：准入检查还够吗？','Revocation during queuing: is admission enough?'),note=L('时间顺序是关键。t0准入有效，t1撤权，t2真正读数据。旧票据不能继续放行。检查与读必须相对于更新共享线性化点。已在撤权前读出的字节不能被自动收回；发布必须是另一个检查。','Time order matters: admission at t0, revocation at t1, read at t2. The old ticket must not authorize a later read. Revalidation and reading need a shared linearization point relative to updates. Bytes released before revocation cannot be retracted; publication is a separate check.'))
    s.line(1.2,3.2,12,3.2,BLUE,3)
    for x,title,body in [(1.2,L('t0 准入','t0 Admission'),L('策略 epoch = 7\n生成票据','Policy epoch = 7\nTicket created')),(5.3,L('t1 撤销授权','t1 Revocation'),L('策略 epoch = 8\n旧票据失效','Policy epoch = 8\nOld ticket invalidated')),(9.4,L('t2 实际执行','t2 Execution'),L('重新检查\n拒绝或重新准入','Revalidate\nReject or re-admit'))]:
        s.rect(x,3.09,.18,.22,TEAL);s.text(x,2.13,3.0,.6,title,22,True);s.text(x,3.75,3,.95,body,19)
    s.takeaway(L('单机模型用锁；分布式适配器仍需实现一致性协议。','The model uses a lock; distributed adapters still need a consistency protocol.'))

    s=S(L('变更影响可沿显式依赖传播','Change impact follows explicit dependencies'),note=L('遍历依赖图本身是成熟算法。值得研究的是是否能完整记录语义、策略、算子、主题与产物的依赖，并减少真实变更重验证工作。当前实现是保守的全局epoch，没有实现细粒度优化。','Dependency traversal is a standard algorithm. The research question is whether semantic, policy, operator, theme and artifact dependencies can be captured completely and reduce real revalidation work. The current model uses a conservative global epoch, not fine-grained optimization.'))
    nodes=[(.9,2.0,L('源／策略变更','Source / policy change')),(4.85,2.0,L('IR 失效','IR invalidated')),(8.8,2.0,L('产物重验证','Artifact revalidation'))]
    for x,y,t in nodes:s.rect(x,y,3.3,.95);s.text(x+.2,y+.2,2.9,.6,t,20,True,color=BLUE)
    s.line(4.2,2.48,4.85,2.48);s.line(8.15,2.48,8.8,2.48)
    s.panel(.9,3.65,5.5,1.85,L('条件：依赖记录完整','Condition: complete dependencies'),L('隐藏视图、共享 prompt、latest 别名都可能破坏局部性','Hidden views, shared prompts and mutable aliases can break locality'),size=17)
    s.panel(6.7,3.65,5.4,1.85,L('现状：保守失效','Current model: conservative'),L('全局 epoch；不宣称细粒度维护成本已经降低','Global epoch; no measured fine-grained maintenance savings'),size=17)

    s=S(L('文件存在，仍不等于任务已完成','An existing file is not a verified deliverable'),note=L('内部团队容易把生成文件当成完成。举例：上次运行留下同名PPT、脚本写好但没运行、引用能打开但不支持结论。这些都应留下未解决义务，而不是勾选完成。','Teams often equate file generation with completion. Examples include a same-named old deck, an unexecuted script, or a valid citation that does not support the claim. These cases retain unresolved obligations instead of being marked complete.'))
    s.panel(.85,1.95,5.4,3.55,L('Theme 完成谓词','Theme completion predicate'),L('所有必需字段齐全\n产物摘要与本轮匹配\n证据依赖仍有效\n指定验证器已通过','All required fields present\nArtifact digest matches this run\nEvidence dependencies remain valid\nSpecified verifier has passed'))
    s.panel(6.65,1.95,5.75,3.55,L('分别记录状态','Keep states distinct'),L('已验证完成\n部分完成\n拒绝\n过期，需要重验证','Completed and verified\nPartial\nRejected\nExpired; revalidation required'),color=BLUE)
    s.takeaway(L('内容哈希证明字节身份；不证明事实为真或语义蕴含。','A content hash proves byte identity, not truth or semantic entailment.'))

    summary=json.loads((ROOT.parent/'agentic-runtime-preprint/artifact_v29/results/summary.json').read_text())
    s=S(L('已经执行：有限契约一致性验证','Executed evidence: finite contract conformance'),
        L('一个收入记录 × 十个二值故障维度；穷举整个声明域','One revenue fixture × ten binary fault dimensions; exhaustive within this domain'),
        [('Raw cases / 逐例结果','../agentic-runtime-preprint/artifact_v29/results/fault-cases.csv'),('Summary / 汇总','../agentic-runtime-preprint/artifact_v29/results/summary.json')],
        L('这些数值从提交的JSON结果读取。1,024种组合中，只有没有故障的组合应通过；其余1,023种应拒绝。零分歧说明这个有限域的规则执行一致。不能把1,024解释为统计样本量，也不能把零分歧换成生产失效率上界。','Numbers are loaded from the committed result JSON. Only the no-fault configuration should pass; the other 1,023 should be rejected. Zero disagreement establishes conformance in this finite domain. It is not a statistical sample or a production failure-rate bound.'))
    for i,(n,label) in enumerate([(summary['cases'],L('声明组合','Declared cases')),(summary['rejected'],L('预期拒绝','Expected rejections')),(summary['oracle_disagreements'],L('与 oracle 分歧','Oracle disagreements'))]):
        x=.9+i*4;s.text(x,2.1,3.5,1.0,f'{n:,}',48,True,color=TEAL);s.text(x,3.24,3.5,.55,label,21)
    s.rect(.9,4.38,11.6,.95);s.text(1.12,4.58,11.1,.53,L('另有 5 项生命周期检查通过：撤权、替换、过期、重读、重验证','Five lifecycle checks pass: revocation, replacement, expiry, replay and revalidation'),19,True)
    s.takeaway(L('这是有限域一致性结果；不是检索质量或扩展性实验。','This is finite conformance, not retrieval-quality or scalability evidence.'))

    s=S(L('证据边界：我们现在可以说什么','Evidence boundary: what can we say today?'),note=L('明确区分实现边界和未来承诺。没有自然检索、并发压力、跨租户攻击或完整P1数据。参考模型依赖可信描述与完整中介，不是安全沙箱。保持这个区分能让工程讨论更有效。','Separate implementation scope from future commitments. There are no natural retrieval, load, cross-tenant attack or complete P1 results. The model assumes trusted descriptors and mediation; it is not a security sandbox.'))
    s.panel(.85,1.95,5.5,3.6,L('可以说','Supported now'),L('数据边界有可运行规格\n十维有限故障域已检查\n五项生命周期检查通过\n代码、逐例结果、哈希可检查','Executable data-boundary specification\nTen-dimensional fault domain checked\nFive lifecycle checks passed\nCode, cases and hashes are inspectable'),size=18)
    s.panel(6.65,1.95,5.8,3.6,L('尚不能说','Not established'),L('提高真实任务正确率\n降低企业总成本\n实现分布式安全或扩展性\nP1 已成立或已达到顶会实证要求','Better real-task accuracy\nLower enterprise lifecycle cost\nDistributed safety or scalability\nP1 support or conference readiness'),color=GOLD,size=18)

    s=S(L('近期工作：执行、研究与交付','Recent work: execution, research and delivery'),refs=SOURCES[:3],note=L('三篇都不是我们的实验结果。Prime Agent说明长期状态与核算已有实现；AutoResearch区分产物和证据；Apodex区分生成和受控交付。v29不能再以这些泛化能力作为独有创新。','These are external studies, not our results. Prime Agent operationalizes persistent state and accounting. AutoResearch separates artifacts from evidence. Apodex separates generation from controlled delivery. Those broad capabilities are not our unique novelty.'))
    for i,(title,body) in enumerate([('Prime Agent',L('持久 REPL、递归会话、全树核算\n迁移：固定总预算和持久状态\n限制：高分不证明 P1','Persistent REPL, recursive sessions, tree-wide accounting\nTransfer: control budget and state\nLimit: scores do not establish P1')),('AutoResearch',L('独立评价与主张—证据绑定\n迁移：结果必须对应真实产物\n限制：示范案例非普适可靠性','Independent review and claim–evidence binding\nTransfer: require real artifacts\nLimit: examples are not universal reliability')),('Apodex 1.1',L('证据综合与受控产物交付\n迁移：交付清单和时效验证\n限制：其两轴与 P1 不同','Evidence synthesis and controlled delivery\nTransfer: manifests and validity\nLimit: its scaling axes differ from P1'))]):
        s.panel(.8+i*4.0,1.95,3.75,3.95,title,body,size=17)

    s=S(L('近期工作：索引、路由与来源读取','Recent work: indexing, routing and source access'),refs=SOURCES[3:],note=L('VoiceMem消融提醒索引与路由要分别控制；NeoHorse记录需求与实际路由，提醒容量变化可能偷偷改变模型；Beyond Top-K要求强BM25对照且不能忽视转换错误。最后一篇是已有引用的复核。','VoiceMem motivates separate index and routing ablations. NeoHorse separates demand from actual route, exposing capacity-driven model changes. Beyond Top-K motivates strong BM25 controls and conversion-error accounting. The last paper is a recheck of an existing citation.'))
    for i,(title,body) in enumerate([('VoiceMem',L('上层索引与后端分开\n迁移：索引、路由分别消融\n限制：检索时延非端到端时延','Separate upper index and backend\nTransfer: isolate index and routing\nLimit: retrieval latency is not end-to-end latency')),('NeoHorse-1',L('预测需求 ≠ 实际执行路由\n迁移：冻结模型并记录改路原因\n限制：一次训练非持续自改进','Predicted demand ≠ served route\nTransfer: freeze models; log rerouting\nLimit: one update is not sustained self-improvement')),('Beyond Top-K',L('来源结构、单位和年份很关键\n迁移：强 BM25 + 原文对照\n限制：单文档小样本','Source structure, units and time matter\nTransfer: strong BM25 + source controls\nLimit: small, single-document study'))]):
        s.panel(.8+i*4.0,1.95,3.75,3.95,title,body,size=17)

    s=S(L('v29 修正了哪些方法学风险','Methodological repairs in v29'),note=L('关键是四态判定不再重叠。先确认是否有足够测量证据，再确认义务是否成立，最后才谈P1。另删除错误的p95功效推算；用pilot模拟决定样本量，不能用噪声扩大允许退化。','The four-state verdicts no longer overlap. Establish measurement adequacy, then obligation eligibility, then classify P1. Remove the invalid p95 power calculation. Pilot simulation determines replication; it must not widen acceptable degradation.'))
    rows=[(L('测量不足','Insufficient measurement'),L('覆盖、校准或区间不能判定 → inconclusive','Unresolved coverage, calibration or intervals → inconclusive')),(L('义务明确失败','Clear obligation violation'),L('充分测量下违规超界 → conditional-engineering','Adequate evidence and violation beyond bound → conditional-engineering')),(L('义务全部通过','All obligations pass'),L('再按语义、交互、成本区间判断 P1','Then classify P1 using semantic, interaction and cost intervals')),(L('阈值与功效','Margins and power'),L('阈值源自决策；p95 功效需要保留尾分布的模拟','Decision-derived margins; p95 power needs tail-aware simulation'))]
    for i,(a,b) in enumerate(rows):
        y=1.9+i*.89;s.rect(.8,y,11.7,.75);s.text(1,y+.15,3.2,.42,a,18,True,color=TEAL);s.text(4.35,y+.14,7.95,.56,b,17)

    s=S(L('下一步实验：隔离契约本身的作用','Next experiment: isolate contract enforcement'),note=L('相同候选、相同工具、相同模型和预算，只改变描述是否被执行边界强制执行。加入自然错误和合成注入，按来源和时间划分留出集。误拒绝必须与违规一起报告，否则拒绝所有请求就会看起来很好。','Hold candidates, tools, model and budget fixed; vary whether the boundary enforces the described constraints. Include natural errors and injections, with source/time holdouts. Report false refusal alongside violations; rejecting everything is not useful.'))
    s.panel(.8,1.92,5.55,3.5,L('中心因果对照','Central causal contrast'),L('A：保留来源的描述式目录\nB：相同目录 + 执行契约\n固定候选、工具、模型、预算\n来源／模板／时间分组留出','A: source-preserving prose catalog\nB: same catalog + enforcement\nFix candidates, tools, model, budget\nHold out sources, templates and time'),size=18)
    s.panel(6.65,1.92,5.8,3.5,L('一起报告','Report together'),L('所有字段正确且证据有效\n未授权／过期释放\n合法任务误拒\n构建 + 维护 + 在线总成本','All fields correct with valid evidence\nUnauthorized / stale release\nFalse refusal on valid tasks\nBuild + maintenance + online cost'),size=18)
    s.takeaway(L('另做强检索基线与逐维消融，避免把导航收益归因于治理。','Add strong retrieval baselines and ablations; separate navigation from governance.'))

    s=S(L('内部试点：先选一个边界清楚的工作流','Internal pilot: start with one bounded workflow'),
        L('建议方案；尚未执行的内部试点','Proposed plan; no internal pilot is claimed'),note=L('建议选季度报告或内部研究简报，从人工可核验字段和一个受控源开始。业务负责人定义可容忍错误，数据负责人定义源与策略，平台负责人实现执行边界。每阶段设进入下一阶段的证据条件，而非承诺任意日期。','Choose a quarterly report or internal research brief with human-verifiable fields and a controlled source. Business owns tolerances; data owners define sources and policy; platform engineers enforce the boundary. Use evidence gates rather than arbitrary schedule promises.'))
    for i,(a,b) in enumerate([(L('阶段 1：定义','Phase 1: specify'),L('选定输出字段与来源\n明确单位、期间、授权\n人工标注正确与拒绝案例','Select output fields and sources\nDefine units, periods and access\nLabel valid and refusal cases')),(L('阶段 2：适配','Phase 2: integrate'),L('接入一个真实读适配器\n实现执行期重验证\n用独立日志核查中介覆盖','Integrate one real read adapter\nImplement execution-time checks\nAudit mediation with independent logs')),(L('阶段 3：比较','Phase 3: evaluate'),L('运行固定预算对照\n报告误拒、质量与总成本\n通过证据门槛后扩展范围','Run fixed-budget comparisons\nReport refusal, quality and total cost\nExpand only after evidence gates pass'))]):
        s.panel(.8+i*4,1.95,3.75,3.7,a,b,size=18)

    sig='https://2027.sigmod.org/calls_papers_sigmod_research.shtml'
    osdi='https://www.usenix.org/conference/osdi27/call-for-papers'
    icse='https://conf.researchr.org/track/icse-2027/icse-2027-research-track'
    s=S(L('投稿方向由主贡献与实证决定','Venue choice follows the contribution and evidence'),
        L('时间核查：2026-09-10；当前稿仍需补实证','Dates checked on 2026-09-10; the current manuscript still needs empirical studies'),
        [('SIGMOD 2027 CFP',sig),('OSDI 2027 CFP',osdi),('ICSE 2027 CFP',icse)],
        L('数据契约路线优先考虑SIGMOD，完整P1运行时路线考虑OSDI，契约演进与维护考虑ICSE/FSE。SIGMOD十月窗口很紧；不能为了赶时间把合成检查当实证。ICSE2027已经截止。FSE2027截止日期未确认。','Prioritize SIGMOD for data contracts, OSDI for an implemented P1 runtime study, and ICSE/FSE for evolution and maintainability. The October SIGMOD window is tight. Synthetic checks cannot replace empirical work. ICSE 2027 is closed; FSE 2027 deadlines were not verified.'))
    rows=[('SIGMOD / PACMMOD',L('数据契约、真实异构数据、维护成本','Data contracts, heterogeneous data, maintenance cost'),L('摘要 10/10；全文 10/17，AoE','Abstract Oct 10; paper Oct 17, AoE')),('OSDI 2027',L('完整运行时、独立监测、规模与故障实验','Implemented runtime, monitoring, scale and failures'),L('摘要 12/1；全文 12/8，22:59 UTC','Abstract Dec 1; paper Dec 8, 22:59 UTC')),('ICSE / FSE',L('契约演进、变更影响、维护实证','Contract evolution, change impact, maintainability'),L('ICSE 2027 已截止；FSE 日期未核实','ICSE 2027 closed; FSE dates unverified'))]
    for i,(a,b,c) in enumerate(rows):
        y=1.95+i*1.19;s.rect(.8,y,11.7,1.03);s.text(1,y+.15,3.4,.45,a,20,True,color=TEAL);s.text(4.3,y+.12,7.95,.42,b,17);s.text(4.3,y+.59,7.95,.34,c,14,color=MUTED)
    s.takeaway(L('先决定论文的主问题，再集中资源补齐对应证据。','Choose the primary research question, then fund the evidence it requires.'))

    s=S(L('内部讨论需要做的三个决定','Three decisions for the internal discussion'),note=L('结尾不要以论文已经完成所有验证作为行动理由。团队需要决定优先研究问题、可用真实数据与合格阈值。产出应是责任明确的试点协议，而不是新的架构命名。','Do not close by implying that all validation is complete. The team must choose the primary question, accessible real data and decision-derived tolerances. The next output is an owned pilot protocol, not another architecture name.'))
    for i,(a,b) in enumerate([(L('选主问题','Choose the question'),L('优先数据使用契约，还是完整能力—容量分离？','Data-use contracts first, or a full capability–capacity study?')),(L('选证据场景','Choose the evidence setting'),L('哪个真实工作流有明确来源、授权和可核验输出？','Which real workflow has clear sources, authority and verifiable outputs?')),(L('定验收条件','Define acceptance criteria'),L('谁决定质量退化、误拒与总成本的可接受范围？','Who owns acceptable quality degradation, false refusal and total cost?'))]):
        y=1.9+i*1.15;s.text(.9,y,3.65,.55,f'0{i+1}  {a}',23,True,color=TEAL);s.text(4.8,y+.02,7.4,.84,b,20)
    s.takeaway(L('v29 的价值：把架构主张变成可执行、可追证、可被反驳的对象。','v29 turns an architectural claim into something executable and falsifiable.'))

    s=S(L('阅读与复现入口','Reading and reproduction'),
        L('五篇新增精读 + 一篇既有文献复核；全部采用原文链接','Five new readings plus one recheck; links point to the original papers'),
        [('Existing arXiv',PAPER),('v29 evidence record','../agentic-runtime-preprint/V29_REVIEW.md')],
        L('页内论文链接可以点击。讲稿备注保留章节级阅读提示。代码和逐例CSV位于同一仓库artifact_v29。现有arXiv链接是前版，本地v29尚未完成arXiv替换。','Paper links on this slide are clickable. Speaker notes preserve reading guidance. The code and case CSV are in artifact_v29 in the same repository. The existing arXiv link points to the earlier public manuscript; v29 has not been submitted as a replacement.'))
    for i,(a,b) in enumerate(SOURCES):
        y=1.83+i*.51;s.text(.9,y,3.35,.37,a,17,True,color=TEAL,link=b);s.text(4.35,y+.01,8,.32,b,13,color=MUTED,link=b)
    s.rect(.9,5.18,11.55,.91);s.text(1.08,5.34,11.15,.59,L('复现：artifact_v29/evaluate.py + test_contract.py\n原始结果：fault-cases.csv / summary.json；P1 未评估','Reproduce: artifact_v29/evaluate.py + test_contract.py\nRaw evidence: fault-cases.csv / summary.json; P1 not evaluated'),16,color=BLUE)

    assert len(slides)==22
    out=ROOT/f'paper-v29-internal-{lang}.pptx';deck.save(out)
    pdf=Canvas(str(ROOT/f'paper-v29-internal-{lang}-preview.pdf'),pagesize=(960,540))
    thumb=Image.new('RGB',(1600,((len(slides)+3)//4)*225),'#DDE3E9')
    notes=[]
    for i,s in enumerate(slides):
        pdf.drawImage(ImageReader(s.image),0,0,width=960,height=540);pdf.showPage()
        thumb.paste(s.image.resize((400,225)),((i%4)*400,(i//4)*225))
        notes.append(f'## {i+1:02d}. {s.texts[0]}\n\n{s.notes}\n\n'+ '\n'.join(f'- [{a}]({b})' for a,b in s.sources))
    pdf.save();thumb.save(ROOT/f'paper-v29-internal-{lang}-contact-sheet.png')
    (ROOT/f'speaker-notes-{lang}.md').write_text('\n\n'.join(notes)+'\n')
    # Reload the serialized deck and verify notes plus basic bounds.
    check=Presentation(out);assert len(check.slides)==22
    for slide in check.slides:
        assert slide.notes_slide.notes_text_frame.text.strip()
        for sh in slide.shapes:
            assert sh.left>=0 and sh.top>=0
            assert sh.left+sh.width<=deck.slide_width+100 and sh.top+sh.height<=deck.slide_height+100
    print(f'{out}: 22 slides, notes and editable shapes verified')


if __name__=='__main__':
    build('zh');build('en')
